from __future__ import annotations

from pathlib import Path


def extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "") if para.style is not None else ""
        if style.startswith("Heading"):
            try:
                level = int(style.split()[-1])
            except ValueError:
                level = 1
            parts.append("#" * min(level, 6) + " " + text)
        else:
            parts.append(text)
    for table in doc.tables:
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        if rows:
            parts.append("\n".join(rows))
    return "\n\n".join(parts)
