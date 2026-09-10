from __future__ import annotations

from pathlib import Path


def extract_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        lines: list[str] = []
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            _shape_text(shape, lines)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        block = [f"--- slide {i}" + (f": {title}" if title else "") + " ---"]
        block.extend(l for l in lines if l)
        if notes:
            block.append(f"[notes] {notes}")
        parts.append("\n".join(block))
    return "\n\n".join(parts)


def _shape_text(shape, out: list[str]) -> None:  # type: ignore[no-untyped-def]
    if shape.shape_type == 6 and hasattr(shape, "shapes"):  # group
        for s in shape.shapes:
            _shape_text(s, out)
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            text = "".join(r.text for r in para.runs).strip()
            if text:
                out.append(("  " * para.level) + text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells]
            out.append(" | ".join(cells))
