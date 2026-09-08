"""Regenerate the tiny binary fixtures used by tests/extract/test_binary.py.

Run: .venv/bin/python tests/fixtures/make_fixtures.py
"""

from pathlib import Path

HERE = Path(__file__).parent


def make_pdf() -> None:
    import pymupdf as fitz

    doc = fitz.open()
    for i, text in enumerate(["Photosynthesis converts light energy.", "Second page about chlorophyll."], 1):
        page = doc.new_page()
        page.insert_text((72, 72), f"Page {i}: {text}")
    doc.save(HERE / "sample.pdf")


def make_pptx() -> None:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Binary Search Trees"
    slide.placeholders[1].text = "Average lookup is O(log n)."
    slide.notes_slide.notes_text_frame.text = "Mention balancing."
    prs.save(HERE / "sample.pptx")


def make_docx() -> None:
    from docx import Document

    d = Document()
    d.add_heading("Essay Guidelines", level=1)
    d.add_paragraph("Submit as PDF by Friday.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Criterion"
    t.rows[0].cells[1].text = "Weight"
    d.save(HERE / "sample.docx")


if __name__ == "__main__":
    make_pdf()
    make_pptx()
    make_docx()
    print("fixtures written to", HERE)
